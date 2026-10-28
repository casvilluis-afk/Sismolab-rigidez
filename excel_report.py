"""Memoria de cálculo en Excel con la misma estructura del trabajo del curso.

Hojas: CARATULA, RIGIDEZ, MASA, ANALISIS SISMICO, CENTRO DE MASA,
CENTRO DE RIGIDEZ, PSEUDOTRIDIMENSIONAL, DESPLAZ LATERALES, MODAL y ESPECTRO.

Unidades MKS: tonf, m, s (masa en tonf·s²/m). Los datos van en azul y los
cálculos son fórmulas de Excel (pesos, masas, centro de masa, cortante basal,
fuerzas por nivel, centro de rigidez, matrices G, Gᵀ, K, K⁻¹ y
desplazamientos), de modo que el libro se recalcula si se cambia un dato.
Las columnas «SismoLab» repiten el valor de la página para comprobar.

No depende del navegador: build_workbook(report) devuelve los bytes .xlsx.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------

NAVY = "1F3864"
BLUE = "2F5597"
HEAD = "D9E1F2"
RESULT = "FFF2CC"
CHECK = "E2EFDA"
INPUT_FONT = Font(name="Calibri", size=10, color="0000FF")
BODY_FONT = Font(name="Calibri", size=10, color="000000")
BOLD = Font(name="Calibri", size=10, bold=True)
WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
NOTE_FONT = Font(name="Calibri", size=9, italic=True, color="595959")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")

F_INT = "0"
F_2 = "#,##0.00"
F_3 = "#,##0.000"
F_4 = "0.0000"
F_6 = "0.000000"
F_SCI = "0.000E+00"


def col(c: int) -> str:
    return get_column_letter(c)


def addr(row: int, column: int, absolute: bool = False) -> str:
    if absolute:
        return f"${col(column)}${row}"
    return f"{col(column)}{row}"


def rng(r1: int, c1: int, r2: int, c2: int, absolute: bool = True) -> str:
    return f"{addr(r1, c1, absolute)}:{addr(r2, c2, absolute)}"


class Sheet:
    def __init__(self, wb: Workbook, title: str, widths: dict[int, float] | None = None):
        self.ws = wb.create_sheet(title)
        self.title = title
        self.ws.sheet_view.showGridLines = False
        for column, width in (widths or {}).items():
            self.ws.column_dimensions[col(column)].width = width

    def ext(self, row: int, column: int) -> str:
        """Referencia absoluta desde otra hoja."""
        return f"'{self.title}'!{addr(row, column, True)}"

    def ext_range(self, r1: int, c1: int, r2: int, c2: int) -> str:
        return f"'{self.title}'!{rng(r1, c1, r2, c2)}"

    def put(self, row, column, value, fmt=None, kind="body", border=True, align=None):
        cell = self.ws.cell(row=row, column=column, value=value)
        if kind == "input":
            cell.font = INPUT_FONT
        elif kind == "head":
            cell.font = BOLD
            cell.fill = PatternFill("solid", fgColor=HEAD)
            cell.alignment = CENTER
        elif kind == "label":
            cell.font = BOLD
        elif kind == "result":
            cell.font = BOLD
            cell.fill = PatternFill("solid", fgColor=RESULT)
        elif kind == "check":
            cell.font = Font(name="Calibri", size=10, color="375623")
            cell.fill = PatternFill("solid", fgColor=CHECK)
        elif kind == "note":
            cell.font = NOTE_FONT
            border = False
        else:
            cell.font = BODY_FONT
        if fmt:
            cell.number_format = fmt
        if border:
            cell.border = BOX
        if align is not None:
            cell.alignment = align
        return cell

    def array(self, r1, c1, r2, c2, formula, fmt=F_3):
        ref = rng(r1, c1, r2, c2, absolute=False)
        self.ws.cell(row=r1, column=c1).value = ArrayFormula(ref, formula)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = self.ws.cell(row=r, column=c)
                cell.number_format = fmt
                cell.border = BOX
                cell.font = BODY_FONT

    def title_bar(self, row, c1, c2, text, size="title"):
        self.ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        cell = self.ws.cell(row=row, column=c1, value=text)
        cell.font = TITLE_FONT if size == "title" else WHITE_BOLD
        cell.fill = PatternFill("solid", fgColor=NAVY if size == "title" else BLUE)
        cell.alignment = LEFT
        self.ws.row_dimensions[row].height = 22 if size == "title" else 18
        for c in range(c1, c2 + 1):
            self.ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=NAVY if size == "title" else BLUE)

    def headers(self, row, c1, values):
        for offset, value in enumerate(values):
            self.put(row, c1 + offset, value, kind="head")

    def note(self, row, column, text):
        self.put(row, column, text, kind="note")


# ---------------------------------------------------------------------------
# Libro
# ---------------------------------------------------------------------------

def build_workbook(report: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    refs: dict[str, object] = {}
    _caratula(wb, report)
    _rigidez(wb, report, refs)
    _masa(wb, report, refs)
    _centro_masa(wb, report, refs)
    _analisis_sismico(wb, report, refs)
    _centro_rigidez(wb, report, refs)
    _pseudo(wb, report, refs)
    _desplazamientos(wb, report, refs)
    _modal(wb, report, refs)
    _espectro(wb, report, refs)
    wb.calculation.fullCalcOnLoad = True
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _caratula(wb, report):
    meta = report["meta"]
    sh = Sheet(wb, "CARATULA", {1: 3, 2: 30, 3: 22, 4: 16, 5: 16, 6: 16})
    sh.title_bar(2, 2, 6, "SISMOLAB · MEMORIA DE ANÁLISIS SÍSMICO")
    rows = [
        ("Curso", "Ingeniería Sismorresistente"),
        ("Proyecto", meta["project"]),
        ("Modelo", meta["model"]),
        ("Uso", meta["use"]),
        ("Sitio", meta["site"]),
        ("Rigidez de pórticos", meta["stiffness_model"]),
        ("Fecha de exportación", meta["date"]),
        ("Unidades", "tonf, m, s · masa en tonf·s²/m · E en tonf/m²"),
    ]
    for index, (label, value) in enumerate(rows):
        sh.put(4 + index, 2, label, kind="label")
        sh.ws.merge_cells(start_row=4 + index, start_column=3, end_row=4 + index, end_column=6)
        sh.put(4 + index, 3, value)
    row = 4 + len(rows) + 1
    sh.put(row, 2, "Docente", kind="label")
    sh.ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
    sh.put(row, 3, "", kind="input")
    row += 2
    sh.headers(row, 2, ["Apellidos y nombres", "Código", "Participación"])
    for offset in range(1, 6):
        for c in range(2, 5):
            sh.put(row + offset, c, None, kind="input")
    row += 7
    sh.title_bar(row, 2, 6, "Contenido", size="section")
    contents = [
        ("RIGIDEZ", "Datos, alturas, inercias y rigidez de cada pórtico por piso"),
        ("MASA", "Metrado por piso: peso, masa y momentos para el centro de masa"),
        ("CENTRO DE MASA", "Resumen por nivel y excentricidad accidental"),
        ("ANALISIS SISMICO", "Z, U, C, S, R, periodo, cortante basal y fuerzas por nivel"),
        ("CENTRO DE RIGIDEZ", "Centro de rigidez por nivel con pórticos en cualquier ángulo"),
        ("PSEUDOTRIDIMENSIONAL", "G, Gᵀ, K del pórtico, Gᵀ·K·G, K global y K⁻¹"),
        ("DESPLAZ LATERALES", "U = K⁻¹·F con torsión accidental y derivas E.030"),
        ("MODAL", "Matriz de masas, periodos, frecuencias y masa participativa"),
        ("ESPECTRO", "Espectro E.030 y análisis modal espectral"),
    ]
    for index, (name, text) in enumerate(contents):
        sh.put(row + 1 + index, 2, name, kind="label")
        sh.ws.merge_cells(start_row=row + 1 + index, start_column=3, end_row=row + 1 + index, end_column=6)
        sh.put(row + 1 + index, 3, text)
    row += len(contents) + 2
    sh.note(row, 2, "Convención: datos en azul · fórmulas en negro · resultados clave en amarillo · verde = valor calculado por SismoLab para comprobar.")
    sh.note(row + 1, 2, meta.get("note", ""))
    sh.note(row + 2, 2, "Uso educativo. Requiere la revisión de un ingeniero estructural.")


def _rigidez(wb, report, refs):
    p = report["params"]
    levels = report["levels"]
    frames = report["frames"]
    n = len(levels)
    sh = Sheet(wb, "RIGIDEZ", {1: 3, 2: 30, 3: 12, 4: 12, 5: 12, 6: 34, 7: 12})
    for c in range(8, 8 + 2 * n + 2):
        sh.ws.column_dimensions[col(c)].width = 12
    sh.title_bar(1, 2, 7 + 2 * n, "RIGIDEZ LATERAL DE LOS PÓRTICOS")
    row = 3
    sh.title_bar(row, 2, 4, "DATOS", size="section")
    row += 1
    sh.headers(row, 2, ["Dato", "Valor", "Unidad"])
    data_rows = []
    if p.get("spans"):
        for name, value in p["spans"].items():
            data_rows.append((name, name, value, "m", F_2))
    data_rows += [
        ("E", "Módulo de elasticidad E", p["E"], "tonf/m²", "#,##0"),
        ("gamma", "Peso específico del concreto γ", p["gamma"], "tonf/m³", F_3),
        ("g", "Gravedad g", p["g"], "m/s²", "0.00000"),
        ("xi", "Amortiguamiento ξ", p["xi"], "", F_2),
    ]
    for key, label, value, unit, fmt in data_rows:
        row += 1
        sh.put(row, 2, label)
        sh.put(row, 3, value, fmt, kind="input")
        sh.put(row, 4, unit)
        refs[f"rig.{key}"] = sh.ext(row, 3)
        refs[f"rig.local.{key}"] = addr(row, 3, True)
    row += 2
    sections = p.get("sections") or []
    if sections:
        sh.title_bar(row, 2, 5, "SECCIONES", size="section")
        row += 1
        sh.headers(row, 2, ["Elemento", "Base b (m)", "Altura h (m)", "Nota"])
        for key, label, b, h, note in sections:
            row += 1
            sh.put(row, 2, label)
            sh.put(row, 3, b, F_2, kind="input")
            sh.put(row, 4, h, F_2, kind="input")
            sh.put(row, 5, note)
            refs[f"sec.{key}.b"] = addr(row, 3, True)
            refs[f"sec.{key}.h"] = addr(row, 4, True)
        row += 2

    rigid = report.get("rigid_formula", False)
    sh.title_bar(row, 2, 6, "ALTURAS DE ENTREPISO", size="section")
    row += 1
    sh.headers(row, 2, ["Piso", "h (m)", "h libre (m)", "H acumulada (m)", "Nota"])
    h_refs, hl_refs, H_refs = [], [], []
    for index, level in enumerate(levels):
        row += 1
        sh.put(row, 2, level["label"])
        sh.put(row, 3, level["height"], F_2, kind="input")
        if rigid and "sec.beam.h" in refs:
            sh.put(row, 4, f"={addr(row, 3)}-{refs['sec.beam.h']}", F_2)
            sh.put(row, 6, "h libre = h − h viga (vigas rígidas)")
        else:
            sh.put(row, 4, f"={addr(row, 3)}", F_2)
            sh.put(row, 6, "Se usa la altura de entrepiso")
        sh.put(row, 5, f"=SUM({addr(row - index, 3, True)}:{addr(row, 3)})", F_2)
        h_refs.append(sh.ext(row, 3))
        hl_refs.append(addr(row, 4, True))
        H_refs.append(sh.ext(row, 5))
    refs["h"] = h_refs
    refs["H"] = H_refs
    row += 2

    sh.title_bar(row, 2, 7 + 2 * n, "RIGIDEZ POR PÓRTICO Y POR PISO", size="section")
    row += 1
    head = ["Pórtico", "β (°)", "Dirección", "x₀ (m)", "y₀ (m)", "I (m⁴)"]
    head += [f"n col. {lv['label']}" for lv in levels]
    head += [f"k {lv['label']} (tonf/m)" for lv in levels]
    sh.headers(row, 2, head)
    first_frame_row = row + 1
    k_refs = []
    for frame in frames:
        row += 1
        sh.put(row, 2, frame["label"])
        sh.put(row, 3, frame["beta"], F_2, kind="input")
        sh.put(row, 4, frame["direction"], align=CENTER)
        sh.put(row, 5, frame["x0"], F_3, kind="input")
        sh.put(row, 6, frame["y0"], F_3, kind="input")
        if rigid and "sec.column.b" in refs:
            bx, by = refs["sec.column.b"], refs["sec.column.h"]
            beta = addr(row, 3)
            sh.put(row, 7, f"=({by}*{bx}^3/12)*COS(RADIANS({beta}))^2+({bx}*{by}^3/12)*SIN(RADIANS({beta}))^2", F_SCI)
        else:
            sh.put(row, 7, frame.get("inertia"), F_SCI, kind="input")
        row_k = []
        for level in range(n):
            n_col = 8 + level
            k_col = 8 + n + level
            sh.put(row, n_col, frame["columns"][level], F_INT, kind="input", align=CENTER)
            if rigid and frame.get("rigid", True):
                ncell = addr(row, n_col)
                formula = f"={ncell}*12*{refs['rig.local.E']}*{addr(row, 7)}/{hl_refs[level]}^3"
                sh.put(row, k_col, formula, F_2)
            else:
                sh.put(row, k_col, frame["story_k"][level], F_2, kind="input")
            row_k.append(sh.ext(row, k_col))
        k_refs.append(row_k)
    last_frame_row = row
    refs["k"] = k_refs
    refs["frame_dir"] = sh.ext_range(first_frame_row, 4, last_frame_row, 4)
    row += 1
    totals = {}
    for direction in ("X", "Y"):
        row += 1
        sh.put(row, 2, f"Σk pórticos en {direction}", kind="label")
        for level in range(n):
            k_col = 8 + n + level
            dir_range = rng(first_frame_row, 4, last_frame_row, 4)
            k_range = rng(first_frame_row, k_col, last_frame_row, k_col)
            sh.put(row, k_col, f'=SUMIF({dir_range},"{direction}",{k_range})', F_2, kind="result")
        totals[direction] = row
    for direction, values in (("X", report["check"].get("sum_k_x")), ("Y", report["check"].get("sum_k_y"))):
        if not values:
            continue
        row += 1
        sh.put(row, 2, f"Σk en {direction} · SismoLab", kind="check")
        for level in range(n):
            sh.put(row, 8 + n + level, values[level], F_2, kind="check")
    row += 2
    if rigid:
        sh.note(row, 2, "k = n·12·E·I / h_libre³ (columnas biempotradas con vigas infinitamente rígidas). I según la orientación β del pórtico.")
    else:
        sh.note(row, 2, "Rigidez de piso equivalente del pórtico (V/Δ con la matriz lateral condensada); los valores vienen de SismoLab.")
    sh.note(row + 1, 2, "Los pórticos inclinados se clasifican por su dirección dominante solo para las sumas; el análisis usa β exacto.")


def _masa(wb, report, refs):
    sh = Sheet(wb, "MASA", {1: 3, 2: 26, 3: 10, 4: 10, 5: 10, 6: 8, 7: 10, 8: 12, 9: 10, 10: 13, 11: 10, 12: 14, 13: 12, 14: 12, 15: 58})
    sh.title_bar(1, 2, 15, "METRADO DE CARGAS, MASA Y CENTRO DE MASA POR PISO")
    g = refs["rig.g"]
    row = 3
    P, M, X, Y = [], [], [], []
    for index, level in enumerate(report["mass"]):
        sh.title_bar(row, 2, 15, level["label"].upper(), size="section")
        row += 1
        sh.headers(row, 2, ["ELEMENTO", "ALTURA (m)", "ANCHO (m)", "LARGO (m)", "CANT.", "P.U", "PESO (tonf)", "g (m/s²)", "MASA (tonf·s²/m)", "X (m)", "Y (m)", "P·X", "P·Y", "NOTA"])
        first = row + 1
        for item in level["rows"]:
            row += 1
            sh.put(row, 2, item["name"])
            sh.put(row, 3, item["h"], F_3, kind="input")
            sh.put(row, 4, item["b"], F_3, kind="input")
            sh.put(row, 5, item["l"], F_3, kind="input")
            sh.put(row, 6, item["n"], "0.00" if item["n"] % 1 else F_INT, kind="input")
            sh.put(row, 7, item["pu"], F_3, kind="input")
            sh.put(row, 8, f"={addr(row, 3)}*{addr(row, 4)}*{addr(row, 5)}*{addr(row, 6)}*{addr(row, 7)}", F_3)
            sh.put(row, 9, f"={g}", "0.00")
            sh.put(row, 10, f"={addr(row, 8)}/{addr(row, 9)}", F_4)
            sh.put(row, 11, item["x"], F_3, kind="input")
            sh.put(row, 12, item["y"], F_3, kind="input")
            sh.put(row, 13, f"={addr(row, 8)}*{addr(row, 11)}", F_3)
            sh.put(row, 14, f"={addr(row, 8)}*{addr(row, 12)}", F_3)
            sh.put(row, 15, item.get("note", ""))
        last = row
        row += 1
        sh.put(row, 7, "SUMA", kind="label")
        for c in (8, 10, 13, 14):
            sh.put(row, c, f"=SUM({rng(first, c, last, c, False)})", F_3 if c != 10 else F_4, kind="result")
        sum_row = row
        row += 1
        sh.put(row, 12, "Xcm, Ycm (m) =", kind="label")
        sh.put(row, 13, f"={addr(sum_row, 13)}/{addr(sum_row, 8)}", F_3, kind="result")
        sh.put(row, 14, f"={addr(sum_row, 14)}/{addr(sum_row, 8)}", F_3, kind="result")
        check = report["check"]["P_tf"][index]
        sh.put(row, 7, "SismoLab", kind="check")
        sh.put(row, 8, check, F_3, kind="check")
        P.append(sh.ext(sum_row, 8))
        M.append(sh.ext(sum_row, 10))
        X.append(sh.ext(row, 13))
        Y.append(sh.ext(row, 14))
        row += 3
    refs["P"], refs["M"], refs["Xcm"], refs["Ycm"] = P, M, X, Y
    sh.note(row, 2, "PESO = ALTURA·ANCHO·LARGO·CANT·P.U. Columnas con altura tributaria (½ entrepiso inferior + ½ superior); vigas con la parte descolgada.")
    sh.note(row + 1, 2, "Losa, acabados, tabiquería y carga viva usan ANCHO = área neta (m²) y P.U en tonf/m²; en la carga viva CANT es el % de la E.030.")


def _centro_masa(wb, report, refs):
    sh = Sheet(wb, "CENTRO DE MASA", {1: 3, 2: 14, 3: 13, 4: 12, 5: 12, 6: 12, 7: 12, 8: 14, 9: 14})
    sh.title_bar(1, 2, 9, "CENTRO DE MASA Y EXCENTRICIDAD ACCIDENTAL")
    row = 3
    sh.headers(row, 2, ["Nivel", "P (tonf)", "Xcm (m)", "Ycm (m)", "Lx (m)", "Ly (m)", "e acc. x = 0,05·Lx", "e acc. y = 0,05·Ly"])
    for index, level in enumerate(report["levels"]):
        row += 1
        sh.put(row, 2, level["label"])
        sh.put(row, 3, f"={refs['P'][index]}", F_3)
        sh.put(row, 4, f"={refs['Xcm'][index]}", F_3, kind="result")
        sh.put(row, 5, f"={refs['Ycm'][index]}", F_3, kind="result")
        sh.put(row, 6, level["plan_x"], F_3, kind="input")
        sh.put(row, 7, level["plan_y"], F_3, kind="input")
        sh.put(row, 8, f"=0.05*{addr(row, 6)}", F_3)
        sh.put(row, 9, f"=0.05*{addr(row, 7)}", F_3)
    row += 2
    a = report["analysis"]
    sh.title_bar(row, 2, 9, "PUNTO DE REFERENCIA DEL ANÁLISIS", size="section")
    row += 1
    for label, value, key in (("X ref (m)", a["cm"][0], "xref"), ("Y ref (m)", a["cm"][1], "yref"), ("e acc. x usada (m)", a["ecc"][0], "ex"), ("e acc. y usada (m)", a["ecc"][1], "ey")):
        sh.put(row, 2, label, kind="label")
        sh.put(row, 3, value, F_4, kind="input")
        refs[f"cm.{key}"] = sh.ext(row, 3)
        row += 1
    sh.note(row, 2, "Las matrices G y el centro de rigidez se refieren a este punto (centro de masa del edificio usado por SismoLab). La excentricidad accidental genera Mt = e·F.")


def _analisis_sismico(wb, report, refs):
    s = report["seismic"]
    levels = report["levels"]
    n = len(levels)
    sh = Sheet(wb, "ANALISIS SISMICO", {1: 3, 2: 34, 3: 14, 4: 36, 5: 3, 6: 12, 7: 10, 8: 10, 9: 13, 10: 14, 11: 10, 12: 13, 13: 14, 14: 14})
    sh.title_bar(1, 2, 14, "ANÁLISIS SÍSMICO ESTÁTICO · E.030:2026")
    row = 3
    sh.title_bar(row, 2, 4, "DATOS DEL PROYECTO", size="section")
    params = [
        ("Z", f"Zona sísmica {s['zone']} · factor Z", s["Z"], None, "input", F_2),
        ("U", f"Categoría {s['category']} · factor de uso U", s["U"], None, "input", F_2),
        ("S", f"Perfil de suelo {s['soil']} · factor S", s["S"], None, "input", F_2),
        ("Tp", "Periodo TP (s)", s["Tp"], None, "input", F_2),
        ("Tl", "Periodo TL (s)", s["Tl"], None, "input", F_2),
        ("R0", f"Sistema: {s['system']} · R₀", s["R0"], None, "input", F_2),
        ("Ia", "Irregularidad en altura Ia", s["Ia"], None, "input", F_2),
        ("Ip", "Irregularidad en planta Ip", s["Ip"], None, "input", F_2),
    ]
    local = {}
    for key, label, value, _f, kind, fmt in params:
        row += 1
        sh.put(row, 2, label)
        sh.put(row, 3, value, fmt, kind=kind)
        local[key] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "R = R₀·Ia·Ip", kind="label")
    sh.put(row, 3, f"={local['R0']}*{local['Ia']}*{local['Ip']}", F_2, kind="result")
    local["R"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "hn: altura total (m)")
    sh.put(row, 3, f"={refs['H'][-1]}", F_2)
    local["hn"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "CT")
    sh.put(row, 3, s["CT"], F_INT, kind="input")
    local["CT"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "T manual (s) · vacío = hn/CT")
    sh.put(row, 3, s.get("period_manual"), F_3, kind="input")
    local["Tm"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "Periodo fundamental T (s)", kind="label")
    sh.put(row, 3, f'=IF(ISNUMBER({local["Tm"]}),{local["Tm"]},{local["hn"]}/{local["CT"]})', F_3, kind="result")
    local["T"] = addr(row, 3, True)
    sh.put(row, 4, "E.030 art. 36: T = hn/CT")
    row += 1
    T, Tp, Tl = local["T"], local["Tp"], local["Tl"]
    sh.put(row, 2, "Factor de amplificación C", kind="label")
    sh.put(row, 3, f"=IF({T}<={Tp},2.5,IF({T}<={Tl},2.5*{Tp}/{T},2.5*{Tp}*{Tl}/{T}^2))", F_3, kind="result")
    sh.put(row, 4, "C = 2,5 hasta TP para la cortante estática (art. 18.3)")
    local["C"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "C usado (C/R ≥ 0,11)")
    sh.put(row, 3, f"=MAX({local['C']},0.11*{local['R']})", F_3)
    local["Cu"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "Coeficiente Z·U·C·S/R", kind="label")
    sh.put(row, 3, f"={local['Z']}*{local['U']}*{local['Cu']}*{local['S']}/{local['R']}", F_4, kind="result")
    local["coef"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "Exponente k")
    sh.put(row, 3, f"=IF({T}<=0.5,1,MIN(0.75+0.5*{T},2))", F_3)
    local["k"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "Peso sísmico total PT (tonf)", kind="label")
    sh.put(row, 3, "=" + "+".join(refs["P"]), F_3, kind="result")
    local["PT"] = addr(row, 3, True)
    row += 1
    sh.put(row, 2, "Cortante basal V = coef·PT (tonf)", kind="label")
    sh.put(row, 3, f"={local['coef']}*{local['PT']}", F_3, kind="result")
    local["V"] = addr(row, 3, True)
    sh.put(row, 4, f"SismoLab: {report['check']['V_tf']:.3f} tonf", kind="check")
    for key in ("Z", "U", "S", "Tp", "Tl", "R", "V", "coef", "T"):
        refs[f"sis.{key}"] = f"'{sh.title}'!{local[key]}"

    top = 3
    sh.title_bar(top, 6, 14, "FUERZAS SÍSMICAS POR NIVEL", size="section")
    sh.headers(top + 1, 6, ["Nivel", "hi (m)", "Hi (m)", "Pi (tonf)", "Pi·Hi^k", "αi", "Fi = αi·V (tonf)", "Vi entrepiso (tonf)", "Fi SismoLab"])
    first = top + 2
    F_refs = [None] * n
    rows_for_level = {}
    for position, level in enumerate(reversed(range(n))):
        r = first + position
        rows_for_level[level] = r
    last = first + n - 1
    for level in range(n):
        r = rows_for_level[level]
        sh.put(r, 6, levels[level]["label"])
        sh.put(r, 7, f"={refs['h'][level]}", F_2)
        sh.put(r, 8, f"={refs['H'][level]}", F_2)
        sh.put(r, 9, f"={refs['P'][level]}", F_3)
        sh.put(r, 10, f"={addr(r, 9)}*{addr(r, 8)}^{local['k']}", F_3)
        sh.put(r, 11, f"={addr(r, 10)}/SUM({rng(first, 10, last, 10)})", F_4)
        sh.put(r, 12, f"={addr(r, 11)}*{local['V']}", F_3, kind="result")
        sh.put(r, 13, f"=SUM({addr(first, 12, True)}:{addr(r, 12)})", F_3)
        sh.put(r, 14, report["check"]["F_tf"][level], F_3, kind="check")
        F_refs[level] = sh.ext(r, 12)
    sh.put(last + 1, 9, f"=SUM({rng(first, 9, last, 9, False)})", F_3, kind="label")
    sh.put(last + 1, 10, f"=SUM({rng(first, 10, last, 10, False)})", F_3, kind="label")
    sh.put(last + 1, 12, f"=SUM({rng(first, 12, last, 12, False)})", F_3, kind="label")
    refs["F"] = F_refs
    sh.note(last + 3, 6, "Fi = αi·V con αi = Pi·Hi^k / Σ Pj·Hj^k (E.030 art. 35). Vi es la cortante del entrepiso i.")


def _centro_rigidez(wb, report, refs):
    frames = report["frames"]
    levels = report["levels"]
    sh = Sheet(wb, "CENTRO DE RIGIDEZ", {1: 3, 2: 16})
    for c in range(3, 16):
        sh.ws.column_dimensions[col(c)].width = 12
    sh.title_bar(1, 2, 15, "CENTRO DE RIGIDEZ POR NIVEL")
    sh.note(2, 2, "Para pórticos en cualquier ángulo β: se busca el punto donde una fuerza de traslación no produce giro. r = (x₀ − Xref)·sen β − (y₀ − Yref)·cos β.")
    xref, yref = refs["cm.xref"], refs["cm.yref"]
    row = 4
    cr_refs = []
    for level_index, level in enumerate(levels):
        sh.title_bar(row, 2, 15, f"CENTRO DE RIGIDEZ · {level['label'].upper()}", size="section")
        row += 1
        sh.headers(row, 2, ["Pórtico", "k (tonf/m)", "β (°)", "cos β", "sen β", "x₀ (m)", "y₀ (m)", "r (m)", "k·cos²", "k·sen²", "k·cos·sen", "k·cos·r", "k·sen·r"])
        first = row + 1
        for index, frame in enumerate(frames):
            row += 1
            sh.put(row, 2, frame["label"])
            sh.put(row, 3, f"={refs['k'][index][level_index]}", F_2)
            sh.put(row, 4, frame["beta"], F_2, kind="input")
            sh.put(row, 5, f"=COS(RADIANS({addr(row, 4)}))", F_4)
            sh.put(row, 6, f"=SIN(RADIANS({addr(row, 4)}))", F_4)
            sh.put(row, 7, frame["x0"], F_3, kind="input")
            sh.put(row, 8, frame["y0"], F_3, kind="input")
            sh.put(row, 9, f"=({addr(row, 7)}-{xref})*{addr(row, 6)}-({addr(row, 8)}-{yref})*{addr(row, 5)}", F_3)
            k, c, s_, r = addr(row, 3), addr(row, 5), addr(row, 6), addr(row, 9)
            sh.put(row, 10, f"={k}*{c}^2", F_2)
            sh.put(row, 11, f"={k}*{s_}^2", F_2)
            sh.put(row, 12, f"={k}*{c}*{s_}", F_2)
            sh.put(row, 13, f"={k}*{c}*{r}", F_2)
            sh.put(row, 14, f"={k}*{s_}*{r}", F_2)
        last = row
        row += 1
        sh.put(row, 9, "Σ", kind="label")
        sums = {}
        for c, key in ((10, "axx"), (11, "ayy"), (12, "axy"), (13, "bx"), (14, "by")):
            sh.put(row, c, f"=SUM({rng(first, c, last, c, False)})", F_2, kind="label")
            sums[key] = addr(row, c)
        row += 1
        det = f"({sums['axx']}*{sums['ayy']}-{sums['axy']}^2)"
        sh.put(row, 2, "X CR (m)", kind="label")
        sh.put(row, 3, f"={xref}+({sums['axx']}*{sums['by']}-{sums['axy']}*{sums['bx']})/{det}", F_3, kind="result")
        sh.put(row, 4, "Y CR (m)", kind="label")
        sh.put(row, 5, f"={yref}+({sums['axy']}*{sums['by']}-{sums['ayy']}*{sums['bx']})/{det}", F_3, kind="result")
        sh.put(row, 6, "e x (m)", kind="label")
        sh.put(row, 7, f"={addr(row, 3)}-{refs['Xcm'][level_index]}", F_3)
        sh.put(row, 8, "e y (m)", kind="label")
        sh.put(row, 9, f"={addr(row, 5)}-{refs['Ycm'][level_index]}", F_3)
        check = report["check"].get("CR")
        if check and level_index == 0:
            sh.put(row, 11, "SismoLab (nivel 1)", kind="check")
            sh.put(row, 12, check[0], F_3, kind="check")
            sh.put(row, 13, check[1], F_3, kind="check")
        cr_refs.append((sh.ext(row, 3), sh.ext(row, 5)))
        row += 3
    refs["CR"] = cr_refs
    sh.note(row - 1, 2, "e = CR − CM del nivel (excentricidad estructural).")


def _pseudo(wb, report, refs):
    frames = report["frames"]
    levels = report["levels"]
    n = len(levels)
    N = 3 * n
    sh = Sheet(wb, "PSEUDOTRIDIMENSIONAL", {1: 3, 2: 12})
    for c in range(3, 12 + 3 * N + 2 * n + 6):
        sh.ws.column_dimensions[col(c)].width = 11
    sh.title_bar(1, 2, 10 + N + n + n + N, "ANÁLISIS PSEUDOTRIDIMENSIONAL · MATRICES DE RIGIDEZ")
    sh.note(2, 2, f"Grados de libertad: Ux1…Ux{n}, Uy1…Uy{n}, θ1…θ{n} (giro antihorario). K_pórtico_global = Gᵀ·K_L·G y K = Σ K_pórtico_global.")
    sh.put(3, 2, "X ref (m)", kind="label")
    sh.put(3, 3, f"={refs['cm.xref']}", F_3)
    sh.put(3, 4, "Y ref (m)", kind="label")
    sh.put(3, 5, f"={refs['cm.yref']}", F_3)
    xref, yref = addr(3, 3, True), addr(3, 5, True)
    dof_labels = [f"Ux{i + 1}" for i in range(n)] + [f"Uy{i + 1}" for i in range(n)] + [f"θ{i + 1}" for i in range(n)]

    c_tab = 2
    c_G = 10
    c_GT = c_G + N + 1
    c_MR = c_GT + n + 1
    c_MRL = c_MR + n + 1
    row = 5
    mrl_blocks = []
    for index, frame in enumerate(frames):
        sh.title_bar(row, 2, c_MRL + N - 1, f"PÓRTICO {frame['label'].upper()}  ·  β = {frame['beta']:.2f}°", size="section")
        row += 1
        sh.headers(row, c_tab, ["Nivel", "x₀ (m)", "y₀ (m)", "α (rad)", "cos α", "sen α", "r (m)"])
        sh.put(row, c_G, f"G{index + 1} (n × 3n)", kind="label", border=False)
        sh.put(row, c_GT, f"G{index + 1}ᵀ", kind="label", border=False)
        sh.put(row, c_MR, f"K_L{index + 1} (tonf/m)", kind="label", border=False)
        sh.put(row, c_MRL, f"G{index + 1}ᵀ·K_L{index + 1}·G{index + 1}", kind="label", border=False)
        top = row + 1
        for level in range(n):
            r = top + level
            sh.put(r, c_tab, levels[level]["label"])
            sh.put(r, c_tab + 1, frame["x0"], F_3, kind="input")
            sh.put(r, c_tab + 2, frame["y0"], F_3, kind="input")
            sh.put(r, c_tab + 3, f"=RADIANS({frame['beta']})", F_6)
            sh.put(r, c_tab + 4, f"=COS({addr(r, c_tab + 3)})", F_6)
            sh.put(r, c_tab + 5, f"=SIN({addr(r, c_tab + 3)})", F_6)
            sh.put(r, c_tab + 6, f"=({addr(r, c_tab + 1)}-{xref})*{addr(r, c_tab + 5)}-({addr(r, c_tab + 2)}-{yref})*{addr(r, c_tab + 4)}", F_4)
            for j in range(N):
                block, lv = divmod(j, n)
                if lv != level:
                    value = 0
                else:
                    value = f"={addr(r, c_tab + 4 + block)}"
                sh.put(r, c_G + j, value, F_4)
        # Gᵀ
        sh.array(top, c_GT, top + N - 1, c_GT + n - 1, f"=TRANSPOSE({rng(top, c_G, top + n - 1, c_G + N - 1)})", F_4)
        # K_L
        rigid = report.get("rigid_formula", False) and frame.get("rigid", True)
        k_refs = refs["k"][index]
        for i in range(n):
            for j in range(n):
                r, c = top + i, c_MR + j
                if rigid:
                    if i == j:
                        value = f"={k_refs[i]}" + (f"+{k_refs[i + 1]}" if i + 1 < n else "")
                    elif j == i + 1:
                        value = f"=-{k_refs[j]}"
                    elif i == j + 1:
                        value = f"=-{k_refs[i]}"
                    else:
                        value = 0
                    sh.put(r, c, value, F_2)
                else:
                    sh.put(r, c, frame["matrix"][i][j], F_2, kind="input")
        gt = rng(top, c_GT, top + N - 1, c_GT + n - 1)
        kl = rng(top, c_MR, top + n - 1, c_MR + n - 1)
        g = rng(top, c_G, top + n - 1, c_G + N - 1)
        sh.array(top, c_MRL, top + N - 1, c_MRL + N - 1, f"=MMULT(MMULT({gt},{kl}),{g})", F_2)
        mrl_blocks.append((top, c_MRL))
        row = top + max(N, n) + 1

    # Matriz global
    sh.title_bar(row, 2, c_MRL + N - 1, "MATRIZ DE RIGIDEZ LATERAL DEL SISTEMA GLOBAL  K = Σ Gᵀ·K_L·G", size="section")
    row += 1
    for j, label in enumerate(dof_labels):
        sh.put(row, c_G + j, label, kind="head")
    top = row + 1
    for i in range(N):
        sh.put(top + i, c_G - 1, dof_labels[i], kind="head")
        for j in range(N):
            terms = "+".join(addr(r0 + i, c0 + j) for r0, c0 in mrl_blocks)
            sh.put(top + i, c_G + j, f"={terms}", F_2)
    K_range = rng(top, c_G, top + N - 1, c_G + N - 1)
    # Control
    check = report["check"].get("K_tf")
    if check:
        c_chk = c_G + N + 2
        sh.put(row, c_chk, "K · SismoLab", kind="check")
        for i in range(N):
            for j in range(N):
                sh.put(top + i, c_chk + j, check[i][j], F_2, kind="check")
    row = top + N + 1
    sh.title_bar(row, 2, c_MRL + N - 1, "MATRIZ DE FLEXIBILIDAD  F = K⁻¹", size="section")
    row += 1
    for j, label in enumerate(dof_labels):
        sh.put(row, c_G + j, label, kind="head")
    inv_top = row + 1
    for i in range(N):
        sh.put(inv_top + i, c_G - 1, dof_labels[i], kind="head")
    sh.array(inv_top, c_G, inv_top + N - 1, c_G + N - 1, f"=MINVERSE({K_range})", F_SCI)
    refs["Kinv"] = sh.ext_range(inv_top, c_G, inv_top + N - 1, c_G + N - 1)
    refs["K"] = sh.ext_range(top, c_G, top + N - 1, c_G + N - 1)


def _desplazamientos(wb, report, refs):
    levels = report["levels"]
    n = len(levels)
    N = 3 * n
    a = report["analysis"]
    sh = Sheet(wb, "DESPLAZ LATERALES", {1: 3, 2: 12, 3: 13, 4: 4, 5: 12, 6: 14, 7: 4, 8: 14, 9: 12, 10: 12, 11: 12, 12: 12, 13: 13, 14: 12, 15: 12})
    sh.title_bar(1, 2, 15, "DESPLAZAMIENTOS LATERALES Y DERIVAS · U = K⁻¹·F")
    row = 3
    sh.put(row, 2, "R", kind="label")
    sh.put(row, 3, f"={refs['sis.R']}", F_2)
    R = addr(row, 3, True)
    sh.put(row, 5, "¿Regular?", kind="label")
    sh.put(row, 6, "SÍ" if a["regular"] else "NO", kind="input", align=CENTER)
    reg = addr(row, 6, True)
    sh.put(row, 8, "Factor 0,75R / 0,85R", kind="label")
    sh.put(row, 9, f'=IF({reg}="SÍ",0.75,0.85)*{R}', F_3, kind="result")
    factor = addr(row, 9, True)
    sh.put(row, 11, "Límite deriva", kind="label")
    sh.put(row, 12, a["drift_limit"], F_4, kind="input")
    limit = addr(row, 12, True)
    row += 1
    sh.put(row, 2, "e acc. x (m)", kind="label")
    sh.put(row, 3, f"={refs['cm.ex']}", F_4)
    ex = addr(row, 3, True)
    sh.put(row, 5, "e acc. y (m)", kind="label")
    sh.put(row, 6, f"={refs['cm.ey']}", F_4)
    ey = addr(row, 6, True)
    sh.note(row, 8, "Mt = ex·Fy − ey·Fx (giro antihorario). Desplazamientos en el punto de referencia (CM).")
    row += 2
    labels_F = [f"Fx{i + 1}" for i in range(n)] + [f"Fy{i + 1}" for i in range(n)] + [f"Mt{i + 1}" for i in range(n)]
    labels_U = [f"Ux{i + 1} (m)" for i in range(n)] + [f"Uy{i + 1} (m)" for i in range(n)] + [f"θ{i + 1} (rad)" for i in range(n)]
    for direction in ("X", "Y"):
        sh.title_bar(row, 2, 15, f"SISMO EN DIRECCIÓN {direction}", size="section")
        row += 1
        sh.headers(row, 2, ["GDL", "F (tonf, tonf·m)"])
        sh.headers(row, 5, ["GDL", "U = K⁻¹·F"])
        sh.put(row, 8, "U · SismoLab", kind="check")
        top = row + 1
        for i in range(N):
            r = top + i
            block, lv = divmod(i, n)
            sh.put(r, 2, labels_F[i], kind="label")
            F = refs["F"][lv]
            if block == 0:
                value = f"={F}" if direction == "X" else 0
            elif block == 1:
                value = f"={F}" if direction == "Y" else 0
            else:
                value = f"=-{ey}*{F}" if direction == "X" else f"={ex}*{F}"
            sh.put(r, 3, value, F_3)
            sh.put(r, 5, labels_U[i], kind="label")
            sh.put(r, 8, report["check"][f"U_{direction}"][i], F_SCI, kind="check")
        sh.array(top, 6, top + N - 1, 6, f"=MMULT({refs['Kinv']},{rng(top, 3, top + N - 1, 3)})", F_SCI)
        row = top + N + 1
        sh.headers(row, 2, ["Piso", "U (m)", "", "Δ (m)", "h (m)", "", "Δ/h elástica", "× factor", "Δ/h inelástica", "Límite", "Estado", "Δ/h ‰"])
        offset = 0 if direction == "X" else n
        first = row + 1
        for level in range(n):
            r = first + level
            u_cell = addr(top + offset + level, 6)
            sh.put(r, 2, levels[level]["label"])
            sh.put(r, 3, f"={u_cell}", F_SCI)
            prev = addr(top + offset + level - 1, 6) if level else None
            sh.put(r, 5, f"={u_cell}-{prev}" if prev else f"={u_cell}", F_SCI)
            sh.put(r, 6, f"={refs['h'][level]}", F_2)
            sh.put(r, 8, f"={addr(r, 5)}/{addr(r, 6)}", F_6)
            sh.put(r, 9, f"={factor}", F_3)
            sh.put(r, 10, f"={addr(r, 8)}*{addr(r, 9)}", F_6, kind="result")
            sh.put(r, 11, f"={limit}", F_4)
            sh.put(r, 12, f'=IF(ABS({addr(r, 10)})<={addr(r, 11)},"CUMPLE","NO CUMPLE")', align=CENTER)
            sh.put(r, 13, f"={addr(r, 10)}*1000", F_2)
        row = first + n + 2
    sh.note(row, 2, "Derivas inelásticas con 0,75R (regular) o 0,85R (irregular), E.030 art. 32; límite 0,007 para concreto armado.")


def _modal(wb, report, refs):
    levels = report["levels"]
    n = len(levels)
    N = 3 * n
    modal = report.get("modal")
    sh = Sheet(wb, "MODAL", {1: 3, 2: 12})
    for c in range(3, 4 + N + 12):
        sh.ws.column_dimensions[col(c)].width = 12
    sh.title_bar(1, 2, 3 + max(N, 12), "ANÁLISIS MODAL · K·φ = ω²·M·φ")
    row = 3
    sh.headers(row, 2, ["Nivel", "m (tonf·s²/m)", "Lx (m)", "Ly (m)", "J = m(Lx²+Ly²)/12", "dx = Xcm − Xref", "dy = Ycm − Yref"])
    m_cells, j_cells, dx_cells, dy_cells = [], [], [], []
    for level in range(n):
        row += 1
        sh.put(row, 2, levels[level]["label"])
        sh.put(row, 3, f"={refs['M'][level]}", F_4)
        sh.put(row, 4, levels[level]["plan_x"], F_3, kind="input")
        sh.put(row, 5, levels[level]["plan_y"], F_3, kind="input")
        sh.put(row, 6, f"={addr(row, 3)}*({addr(row, 4)}^2+{addr(row, 5)}^2)/12", F_3)
        sh.put(row, 7, f"={refs['Xcm'][level]}-{refs['cm.xref']}", F_4)
        sh.put(row, 8, f"={refs['Ycm'][level]}-{refs['cm.yref']}", F_4)
        m_cells.append(addr(row, 3, True))
        j_cells.append(addr(row, 6, True))
        dx_cells.append(addr(row, 7, True))
        dy_cells.append(addr(row, 8, True))
    row += 2
    sh.title_bar(row, 2, 3 + N, "MATRIZ DE MASAS M (orden Ux, Uy, θ)", size="section")
    row += 1
    top = row
    for i in range(N):
        bi, li = divmod(i, n)
        for j in range(N):
            bj, lj = divmod(j, n)
            value = 0
            if li == lj:
                m, J, dx, dy = m_cells[li], j_cells[li], dx_cells[li], dy_cells[li]
                pair = (bi, bj)
                if pair in ((0, 0), (1, 1)):
                    value = f"={m}"
                elif pair in ((0, 2), (2, 0)):
                    value = f"=-{m}*{dy}"
                elif pair in ((1, 2), (2, 1)):
                    value = f"={m}*{dx}"
                elif pair == (2, 2):
                    value = f"={J}+{m}*({dx}^2+{dy}^2)"
            sh.put(top + i, 3 + j, value, F_3)
    row = top + N + 1
    if not modal:
        sh.note(row, 2, "Sin resultados modales: completa el metrado y el análisis en SismoLab.")
        return
    sh.title_bar(row, 2, 14, "PERIODOS, FRECUENCIAS Y MASA PARTICIPATIVA (SismoLab)", size="section")
    row += 1
    sh.headers(row, 2, ["Modo", "T (s)", "f (Hz)", "ω (rad/s)", "Γx", "Γy", "Γθ", "%Mx", "%My", "%Mθ", "Σ%Mx", "Σ%My", "Dominante"])
    for index, mode in enumerate(modal["modes"]):
        row += 1
        values = [index + 1, mode["T"], mode["f"], mode["omega"], mode["gamma"]["X"], mode["gamma"]["Y"], mode["gamma"]["RZ"],
                  mode["ratio"]["X"], mode["ratio"]["Y"], mode["ratio"]["RZ"], mode["cum"]["X"], mode["cum"]["Y"], mode["dominant"]]
        fmts = [F_INT, F_4, F_4, F_3, F_4, F_4, F_4, "0.00%", "0.00%", "0.00%", "0.00%", "0.00%", None]
        for offset, (value, fmt) in enumerate(zip(values, fmts)):
            sh.put(row, 2 + offset, value, fmt, kind="result" if offset in (1, 11, 12) and index == 0 else "body")
    row += 2
    sh.note(row, 2, "Γ y % de masa con φ normalizada a desplazamiento máximo 1. La E.030 (art. 29.1.2) pide considerar los modos que sumen al menos 90 % de la masa.")
    row += 2
    sh.title_bar(row, 2, 3 + len(modal["modes"]), "FORMAS DE MODO φ (traslación máxima = 1)", size="section")
    row += 1
    sh.put(row, 2, "GDL", kind="head")
    for index in range(len(modal["modes"])):
        sh.put(row, 3 + index, f"Modo {index + 1}", kind="head")
    labels = [f"Ux{i + 1}" for i in range(n)] + [f"Uy{i + 1}" for i in range(n)] + [f"θ{i + 1}" for i in range(n)]
    for i in range(N):
        sh.put(row + 1 + i, 2, labels[i], kind="label")
        for index, mode in enumerate(modal["modes"]):
            sh.put(row + 1 + i, 3 + index, mode["shape"][i], F_SCI)


def _espectro(wb, report, refs):
    sh = Sheet(wb, "ESPECTRO", {1: 3, 2: 10, 3: 10, 4: 12, 5: 12, 6: 4})
    for c in range(7, 22):
        sh.ws.column_dimensions[col(c)].width = 12
    sh.title_bar(1, 2, 16, "ESPECTRO DE DISEÑO E.030 Y ANÁLISIS MODAL ESPECTRAL")
    row = 3
    sh.headers(row, 2, ["T (s)", "C", "Sa/g = ZUCS/R", "Sa (m/s²)"])
    Z, U, S, R, Tp, Tl = (refs[f"sis.{k}"] for k in ("Z", "U", "S", "R", "Tp", "Tl"))
    g = refs["rig.g"]
    periods = [0.0, 0.02, 0.05] + [round(0.1 * i, 2) for i in range(1, 41)]
    for T in periods:
        row += 1
        sh.put(row, 2, T, F_2, kind="input")
        t = addr(row, 2)
        sh.put(row, 3, f"=IF({t}<0.2*{Tp},1+7.5*{t}/{Tp},IF({t}<={Tp},2.5,IF({t}<={Tl},2.5*{Tp}/{t},2.5*{Tp}*{Tl}/{t}^2)))", F_3)
        sh.put(row, 4, f"={Z}*{U}*{addr(row, 3)}*{S}/{R}", F_4)
        sh.put(row, 5, f"={addr(row, 4)}*{g}", F_3)
    sh.note(row + 1, 2, "Tabla N° 6: C = 1 + 7,5·T/TP para T < 0,2·TP (análisis dinámico).")

    spectral = report.get("spectral")
    c0 = 7
    r = 3
    if spectral:
        source = "E.030 de diseño (ZUCS/R)" if spectral["source"] == "e030" else "registro cargado ÷ R"
        rule = "CQC (ξ = 5 %)" if spectral["combination"] == "cqc" else "0,25Σ|r| + 0,75√Σr²"
        sh.put(r, c0, f"Espectro: {source} · combinación {rule} · desplazamientos × {spectral['amplification']:.3f}", kind="label", border=False)
        r += 2
        for item in spectral["directions"]:
            d = item["direction"]
            sh.title_bar(r, c0, c0 + 7, f"SISMO EN {d}", size="section")
            r += 1
            sh.headers(r, c0, ["Modo", "T (s)", "Sa (m/s²)", "Γ", "%M", "V basal (tonf)"])
            for index, mode in enumerate(item["modes"]):
                r += 1
                for offset, (value, fmt) in enumerate(zip([index + 1, mode["period"], mode["sa"], mode["gamma"], mode["ratio"], mode["base_shear"]], [F_INT, F_4, F_4, F_4, "0.00%", F_3])):
                    sh.put(r, c0 + offset, value, fmt)
            r += 1
            sh.put(r, c0, "V dinámica (tonf)", kind="label")
            sh.put(r, c0 + 2, item["v_dyn"], F_3, kind="result")
            sh.put(r, c0 + 3, "V estática", kind="label")
            sh.put(r, c0 + 4, item["v_static"], F_3)
            sh.put(r, c0 + 5, "Factor de escala", kind="label")
            sh.put(r, c0 + 6, item["scale"], F_3, kind="result")
            r += 2
            sh.headers(r, c0, ["Nivel", "V piso escalada (tonf)", "u elástico (m)", "u inelástico (m)", "Deriva elástica", "Deriva inelástica", "Estado"])
            for level in reversed(item["levels"]):
                r += 1
                values = [level["label"], level["shear"], level["u"], level["u_inelastic"], level["drift"], level["drift_inelastic"], "CUMPLE" if level["ok"] else "NO CUMPLE"]
                for offset, (value, fmt) in enumerate(zip(values, [None, F_3, F_SCI, F_SCI, F_6, F_6, None])):
                    sh.put(r, c0 + offset, value, fmt)
            r += 3
    else:
        sh.note(r, c0, "Sin análisis modal espectral (faltan modos o datos del metrado).")
        r += 2
    record = report.get("record_spectra")
    if record:
        names = list(record)
        sh.title_bar(r, c0, c0 + len(names), "ESPECTROS DEL REGISTRO · PSa (m/s²)", size="section")
        r += 1
        sh.headers(r, c0, ["T (s)"] + names)
        rows = record[names[0]]
        for index, item in enumerate(rows):
            r += 1
            sh.put(r, c0, item["T"], F_2)
            for offset, name in enumerate(names):
                sh.put(r, c0 + 1 + offset, record[name][index]["PSa"], F_4)
